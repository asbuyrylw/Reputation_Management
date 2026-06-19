"""
Reputation Crowding-Out Engine -- Module 4 (v2): Client Report Generator
=======================================================================
The renewal driver. Generates a polished, persuasive monthly report (.docx) from
everything in Postgres: trend charts across all audit runs, BEFORE/AFTER examples
of how AI answers actually changed, work-order execution status, the correlational
attribution narrative, cost/COGS (internal), and a non-technical executive summary.

Upgrades over v1:
  - Executive summary written for a non-technical owner.
  - Trend charts (matplotlib) across ALL runs, embedded as images.
  - Before/after answer snippets -- the single most persuasive element.
  - Work-order execution status (from the tracking module).
  - Attribution narrative (assets shipped -> metric movement, labeled correlational).

Run:
    python -m rep_engine.report_generator monthly --business-id 1
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import tempfile
from datetime import date


try:
    from .db import db
except ImportError:  # pragma: no cover
    from db import db  # type: ignore

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s: %(message)s")
log = logging.getLogger("report_generator")

# Absolute at import so the generating process and the API agree on the location regardless
# of each one's working directory (the download endpoint anchors served files to this dir).
OUTPUT_DIR = os.path.abspath(os.getenv("REP_OUTPUT_DIR", "output"))                       # PH 2

NAVY = "1F3A5F"
GOLD = "B08D2E"
GREEN = "2E6B3E"
RUST = "8A3B2E"




def _run_series(conn, business_id: int) -> list:
    runs = conn.execute(
        "SELECT id, finished_at FROM audit_runs WHERE business_id=%s AND finished_at IS NOT NULL "
        "ORDER BY id ASC", (business_id,),
    ).fetchall()
    series = []
    for r in runs:
        m = conn.execute(
            "SELECT AVG(goal_alignment) ga, "
            "AVG(CASE WHEN mentions_contested THEN 1 ELSE 0 END) contested, "
            "AVG(CASE WHEN surfaces_owned THEN 1 ELSE 0 END) owned "
            "FROM answers WHERE run_id=%s AND NOT COALESCE(failed,false)", (r["id"],),
        ).fetchone()
        series.append({
            "run_id": r["id"],
            "date": r["finished_at"].date().isoformat() if r["finished_at"] else "",
            "goal_alignment": float(m["ga"]) if m["ga"] is not None else None,
            "contested_rate": float(m["contested"]) if m["contested"] is not None else None,
            "owned_rate": float(m["owned"]) if m["owned"] is not None else None,
        })
    return series


def _before_after(conn, business_id: int, limit: int = 4) -> list:
    runs = conn.execute(
        "SELECT id FROM audit_runs WHERE business_id=%s AND finished_at IS NOT NULL ORDER BY id",
        (business_id,),
    ).fetchall()
    if len(runs) < 2:
        return []
    first, last = runs[0]["id"], runs[-1]["id"]

    def best_answer(run_id, prompt, engine):
        return conn.execute(
            "SELECT engine, answer_text, goal_alignment, mentions_contested FROM answers "
            "WHERE run_id=%s AND prompt=%s AND engine=%s AND answer_text <> '' "
            "ORDER BY goal_alignment DESC NULLS LAST LIMIT 1", (run_id, prompt, engine),
        ).fetchone()

    def common_engine(prompt):
        # Pick ONE engine that answered this prompt in BOTH the first and last run so
        # Before/Now is an honest like-for-like comparison (the old code took the
        # max-goal_alignment answer per run regardless of engine, so "Before" could be
        # Perplexity while "Now" was ChatGPT). Prefer the engine with the strongest
        # latest answer among those present in both runs.
        row = conn.execute(
            "SELECT a1.engine FROM answers a0 JOIN answers a1 USING (engine) "
            "WHERE a0.run_id=%s AND a1.run_id=%s AND a0.prompt=%s AND a1.prompt=%s "
            "AND a0.answer_text <> '' AND a1.answer_text <> '' "
            "ORDER BY a1.goal_alignment DESC NULLS LAST LIMIT 1",
            (first, last, prompt, prompt),
        ).fetchone()
        return row["engine"] if row else None

    prompts = conn.execute(
        "SELECT DISTINCT prompt FROM answers WHERE run_id=%s AND answer_text <> ''", (first,),
    ).fetchall()
    pairs = []
    for p in prompts:
        prompt = p["prompt"]
        eng = common_engine(prompt)
        if not eng:
            continue  # no engine answered this prompt in both runs -> no honest pair
        a0, a1 = best_answer(first, prompt, eng), best_answer(last, prompt, eng)
        if a0 and a1:
            improvement = ((a1["goal_alignment"] or 0) - (a0["goal_alignment"] or 0))
            pairs.append({"prompt": prompt, "engine": eng, "before": a0, "after": a1,
                          "improvement": improvement})
    pairs.sort(key=lambda x: x["improvement"], reverse=True)
    return pairs[:limit]


def _load(business_id: int) -> dict:
    with db() as conn:
        b = conn.execute("SELECT * FROM businesses WHERE id=%s", (business_id,)).fetchone()
        if not b:
            raise SystemExit(f"No business id {business_id}")
        gap_row = conn.execute(
            "SELECT model FROM gap_models WHERE business_id=%s ORDER BY id DESC LIMIT 1",
            (business_id,),
        ).fetchone()
        plan_row = conn.execute(
            "SELECT plan FROM strategy_plans WHERE business_id=%s ORDER BY id DESC LIMIT 1",
            (business_id,),
        ).fetchone()
        series = _run_series(conn, business_id)
        ba = _before_after(conn, business_id)
        wo_counts, assets_n, month_cost = {}, 0, 0.0
        try:
            rows = conn.execute(
                "SELECT status, COUNT(*) n FROM work_orders WHERE business_id=%s GROUP BY status",
                (business_id,),
            ).fetchall()
            wo_counts = {r["status"]: r["n"] for r in rows}
            assets_n = conn.execute(
                "SELECT COUNT(*) n FROM assets WHERE business_id=%s", (business_id,)
            ).fetchone()["n"]
        except Exception:
            pass
        try:
            month_cost = float(conn.execute(
                "SELECT COALESCE(SUM(est_cost_usd),0) s FROM cost_ledger "
                "WHERE business_id=%s AND created_at >= date_trunc('month', now())",
                (business_id,),
            ).fetchone()["s"])
        except Exception:
            pass
        attr = conn.execute(
            "SELECT metric, delta, assets_in_window FROM attribution "
            "WHERE business_id=%s ORDER BY id DESC LIMIT 3", (business_id,),
        ).fetchall()
    return {
        "business": dict(b),
        "gap": (gap_row["model"] if gap_row and isinstance(gap_row["model"], dict)
                else (json.loads(gap_row["model"]) if gap_row else {})),
        "plan": (plan_row["plan"] if plan_row and isinstance(plan_row["plan"], dict)
                 else (json.loads(plan_row["plan"]) if plan_row else {})),
        "series": series, "before_after": ba, "wo_counts": wo_counts,
        "assets_n": assets_n, "month_cost": month_cost,
        "attribution": [dict(a) for a in attr],
    }


def _trend_chart(series: list):
    if len(series) < 2:
        return None
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    dates = [s["date"] for s in series]
    fig, ax = plt.subplots(figsize=(6.5, 3.2), dpi=150)
    ax.plot(dates, [s["goal_alignment"] for s in series], marker="o", label="Goal alignment", color="#2E6B3E", linewidth=2)
    ax.plot(dates, [s["owned_rate"] for s in series], marker="s", label="Owned-content surfacing", color="#1F3A5F", linewidth=2)
    ax.plot(dates, [s["contested_rate"] for s in series], marker="^", label="Contested-term mentions", color="#8A3B2E", linewidth=2)
    ax.set_ylabel("Rate / score")
    ax.set_title("AI Visibility Trend Across Audits")
    ax.legend(fontsize=7, loc="best")
    ax.grid(True, alpha=0.3)
    fig.autofmt_xdate(rotation=30)
    fig.tight_layout()
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    tmp.close()  # close our handle first: Windows can't unlink a still-open temp file
    fig.savefig(tmp.name)
    plt.close(fig)
    return tmp.name


def _doc_helpers(doc):
    """The heading / body / bullet paragraph helpers, bound to `doc`. Extracted
    from generate() unchanged so sections can be built by standalone helpers."""
    from docx.shared import Pt, RGBColor

    def heading(text, size=15, color=NAVY, space_before=12):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color)
        return p

    def body(text, italic=False, color=None, bold=False):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text); r.italic = italic; r.bold = bold
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return p

    def bullet(text):
        doc.add_paragraph(text, style="List Bullet")

    return heading, body, bullet


def _section_timeline(heading, body, business_id):
    # ---- Projected timeline (range-based; clearly a projection) ----
    try:
        heading("Projected Timeline")
        from . import timeline_estimator as _te
        proj = _te.estimate(business_id, quiet=True)
        p = proj["projection"]
        body(f"Estimated time for the accurate narrative to dominate (confidence: "
             f"{proj['confidence']}): expected ~{p['expected']['months']} months "
             f"(range {p['optimistic']['months']}–{p['conservative']['months']} months).")
        body("This is a projection, not a guarantee — it reflects how much accurate "
             "content is being shipped and this business's own measured rate of change, "
             "which sharpens the estimate as more audits accumulate.", italic=True)
    except Exception as e:  # noqa: BLE001  -- never let the projection break the report
        log.warning("report: timeline section unavailable (%s)", e)
        body("A timeline projection is not available for this reporting period.", italic=True)


def _section_acceleration(heading, body, business_id):
    # ---- Ways to accelerate (third-party / human levers) ----
    try:
        heading("Ways to Accelerate")
        from . import acceleration_advisor as _aa
        adv = _aa.advise(business_id, quiet=True)
        body("These are third-party / human actions (which we can't automate for you) "
             "that would compress the expected window. Quantities are per month.")
        for lv in adv["levers_ranked_by_impact"][:4]:
            sm = lv["suggested_per_month"]; ws = lv["weeks_saved_range"]
            body(f"• {sm['low']}–{sm['high']} {lv['unit']}: est. {ws['low']}–{ws['high']} weeks faster. {lv['note']}")
        agg = adv["scenarios"]["aggressive_lift"]["window"]
        mod = adv["scenarios"]["moderate_lift"]["window"]
        body(f"Bundled: a moderate monthly lift lands around ~{mod['months']} months; "
             f"an aggressive lift around ~{agg['months']} months (vs ~"
             f"{adv['baseline_expected_window']['months']} on the current plan alone).")
        body("Reviews and placements must be genuine and within each platform's rules — "
             "never fabricated or incentivized against policy.", italic=True)
    except Exception as e:  # noqa: BLE001
        log.warning("report: acceleration section unavailable (%s)", e)
        body("Acceleration levers are not available for this reporting period.", italic=True)


def _section_share_of_voice(heading, body, business_id):
    # ---- Share of voice (which sources AI engines cite) ----
    try:
        from . import citation_analytics as _ca
        sov = _ca.analyze(business_id, quiet=True)
        if sov and sov.get("total_citations"):
            heading("Share of Voice (AI Citations)")
            parts = ", ".join(f"{k} {round(v*100)}%" for k, v in sov["share_of_voice"].items())
            body(f"Across {sov['total_citations']} citations AI engines made about you, "
                 f"the split was: {parts}.")
            body("The aim is to grow the owned + neutral share so accurate sources dominate "
                 "what AI surfaces — not to remove contested ones.", italic=True)
            mom = _ca.momentum(business_id, quiet=True)
            if mom and mom.get("summary"):
                gain = mom["summary"].get("owned_or_neutral_gainers", [])
                loss = mom["summary"].get("contested_losers", [])
                if gain:
                    body(f"Gaining ground: {', '.join(gain[:5])}.")
                if loss:
                    body(f"Contested sources losing ground: {', '.join(loss[:5])}.")
    except Exception as e:  # noqa: BLE001
        log.warning("report: share-of-voice section unavailable (%s)", e)


def _section_per_engine(heading, body, bullet, business_id):
    # ---- Per-engine coverage & grounding: the cross-engine promise, reported honestly ----
    try:
        from . import ai_state_audit as _ai
        pe = _ai.per_engine_metrics(business_id)
        engines = pe.get("engines") or {}
        if not engines:
            return
        heading("What Each AI Engine Says (and How Grounded)")
        cov = pe.get("coverage") or {}
        if cov.get("partial"):
            cfg = cov.get("configured") or []
            exp = cov.get("expected") or []
            body("Partial coverage: this audit covered "
                 f"{', '.join(cfg) or 'no engines'} ({len(cfg)} of {len(exp)} AI "
                 f"engines). Configure the missing engines ({', '.join(cov.get('missing') or [])}) "
                 "for a complete cross-engine read.", italic=True)
        for name, m in engines.items():
            ga = m.get("goal_alignment") or {}
            gr = m.get("grounded_rate")
            n = m.get("n", 0)
            ga_txt = "n/a" if ga.get("mean") is None else f"{ga['mean']:+.2f}"
            if ga.get("low") is not None:
                ga_txt += f" (95% CI {ga['low']:+.2f}..{ga['high']:+.2f})"
            grounded_txt = ("grounding not reported" if not gr
                            else f"{round(gr['p'] * 100)}% grounded in live web (n={gr['n']})")
            bullet(f"{name}: goal alignment {ga_txt} over {n} answers; {grounded_txt}.")
        body("Goal alignment runs -1 (works against the goal) to +1 (strongly supports it); "
             "ranges are 95% confidence intervals. “Grounded” means the engine answered "
             "from a live web search rather than model memory — ungrounded answers reflect what "
             "the model already believed, not the current web.", italic=True)
    except Exception as e:  # noqa: BLE001
        log.warning("report: per-engine section unavailable (%s)", e)


def _section_competitor(heading, body, bullet, business_id):
    # ---- Competitor benchmarking (only if a benchmark has been run) ----
    try:
        from . import competitor as _comp
        cmp = _comp.compare(business_id, quiet=True)
        if cmp and cmp.get("prompts_compared"):
            heading("How You Compare (AI Share of Voice vs. Competitors)")
            rank = cmp.get("subject_rank")
            field = cmp.get("field_size")
            body(f"Across {cmp['prompts_compared']} category questions, here's how often "
                 f"each name shows up in AI answers. You currently rank #{rank} of {field}.")
            for s in cmp.get("standings", []):
                marker = "  ← you" if s["is_subject"] else ""
                bullet(f"{s['name']}: appears in {round(s['appearance_rate']*100)}% of "
                       f"category questions{marker}")
            # one head-to-head line for the strongest rival, if any gap exists
            h2h = cmp.get("head_to_head", [])
            gap_lines = [h for h in h2h if h.get("competitor_only_prompts", 0) > 0]
            if gap_lines:
                top = max(gap_lines, key=lambda h: h["competitor_only_prompts"])
                body(f"Biggest opportunity: {top['competitor']} appears in "
                     f"{top['competitor_only_prompts']} questions where you don't yet — "
                     f"those are the queries to target with accurate, owned content.", italic=True)
            body("Appearance rate = share of category questions where the name is "
                 "mentioned or its site is cited (a transparent heuristic, not a "
                 "judgment of quality).", italic=True)
    except Exception as e:  # noqa: BLE001
        log.warning("report: competitor section unavailable (%s)", e)


def _save_report(doc, b) -> str:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    safe = "".join(c for c in b.get("name", "business") if c.isalnum() or c in " -_").strip().replace(" ", "_")
    filename = f"{safe}_AI_Visibility_Report_{date.today().isoformat()}.docx"
    path = os.path.join(OUTPUT_DIR, filename)
    doc.save(path)
    _fix_settings_zoom(path)
    log.info("Report written: %s", path)
    print(path)
    _record_report(b.get("id"), filename, os.path.abspath(path))
    return path


def _record_report(business_id, filename: str, path: str) -> None:
    """Track the saved report so the console can list + download it. Best-effort: never let
    a bookkeeping miss (e.g. pre-migration DB) fail an otherwise-good report generation."""
    if not business_id:
        return
    try:
        with db() as conn:
            conn.execute(
                "INSERT INTO reports (business_id, filename, path, kind) VALUES (%s,%s,%s,'monthly')",
                (business_id, filename, path),
            )
            conn.commit()
    except Exception as e:  # noqa: BLE001 -- bookkeeping must not break report generation
        log.warning("report: could not record report row (%s)", e)


def _section_root_cause_and_incidents(heading, body, bullet, business_id):
    """Surface the agentic layer's output -- Graph 1 root-cause + Graph 4 pending
    incidents -- when present. Read-only and a NO-OP when those tables are empty or
    absent (so the report is unchanged unless the agents have run)."""
    try:
        with db() as conn:
            rc = conn.execute("SELECT model FROM root_cause WHERE business_id=%s "
                              "ORDER BY id DESC LIMIT 1", (business_id,)).fetchone()
            inc = conn.execute("SELECT severity, COUNT(*) n FROM incidents WHERE business_id=%s "
                               "AND status='pending_human_review' GROUP BY severity",
                               (business_id,)).fetchall()
    except Exception as e:  # noqa: BLE001 -- tables may be absent on an un-migrated DB
        log.warning("report: agentic-findings section unavailable (%s)", e)
        return
    model = (rc.get("model") if rc else None) or {}
    if not isinstance(model, dict):       # tolerate a hand-edited / schema-drifted row
        model = {}
    summary = str(model.get("summary") or "")
    if summary and "No contested sources" not in summary:
        heading("Why the Contested Narrative Surfaces")
        body(summary[:1500])              # bound a pathological injected string
        sources = model.get("primary_sources")
        for s in (sources if isinstance(sources, list) else [])[:3]:
            if isinstance(s, dict):
                bullet(f"{s.get('url', '')} — {s.get('why', '')}")
        counters = model.get("recommended_counters")
        counters = counters if isinstance(counters, list) else []
        if counters:
            body("Recommended counters: " + "; ".join(str(c) for c in counters[:5]) + ".", italic=True)
    if inc:
        heading("New Contested Mentions Flagged This Period")
        for r in inc:
            bullet(f"{r['severity']}: {r['n']} awaiting review")
        body("These were auto-triaged with a drafted response pending your approval.", italic=True)


def _render_brief(body, bullet, row):
    """Render one production brief as a compact, human-actionable card. Shape-robust:
    a hand-edited / schema-drifted JSONB row never raises (bounded strings, isinstance
    guards) -- same hardening as _section_root_cause_and_incidents."""
    brief = row.get("brief") if isinstance(row.get("brief"), dict) else {}
    title = str(row.get("title") or brief.get("title") or "(untitled)")[:160]
    platform = str(row.get("platform") or brief.get("platform") or "")[:40]
    fmt = str(brief.get("format") or "")[:60]
    length = brief.get("target_length") or brief.get("target_length_seconds")
    length = str(length) if length not in (None, "") else ""
    if length and row.get("channel") == "video" and length.isdigit():
        length = f"{length}s"
    meta = " · ".join(x for x in (platform, fmt, length) if x)
    body(title + (f"  ({meta})" if meta else ""), bold=True)
    tq = str(row.get("target_query") or brief.get("target_query") or "")
    if tq:
        bullet(f"Targets the query: {tq[:200]}")
    kws = brief.get("keywords")
    kws = kws if isinstance(kws, list) else []
    if kws:
        bullet("Keywords: " + ", ".join(str(k) for k in kws[:12])[:300])
    hook = str(brief.get("hook") or "")
    if hook:
        bullet(f"Hook: {hook[:200]}")
    outline = brief.get("outline")
    outline = outline if isinstance(outline, list) else []
    if outline:
        bullet("Outline: " + " → ".join(str(o) for o in outline[:8])[:400])
    cta = str(brief.get("cta") or "")
    if cta:
        bullet(f"Call to action: {cta[:200]}")


def _section_production_briefs(heading, body, bullet, business_id):
    """Surface the video + social PRODUCTION BRIEFS (production_brief.plan) -- the
    specs for content the client/team produces OFF-platform. Read-only and a NO-OP
    when there are no open briefs (so the report is unchanged unless that step ran)."""
    try:
        with db() as conn:
            rows = conn.execute(
                "SELECT channel, platform, title, target_query, brief FROM production_briefs "
                "WHERE business_id=%s AND status='to_produce' ORDER BY channel, id",
                (business_id,)).fetchall()
    except Exception as e:  # noqa: BLE001 -- table may be absent on an un-migrated DB
        log.warning("report: production-briefs section unavailable (%s)", e)
        return
    if not rows:
        return
    vids = [r for r in rows if r.get("channel") == "video"][:8]
    socs = [r for r in rows if r.get("channel") == "social"][:8]
    if not vids and not socs:
        return
    heading("Content to Produce This Period")
    body("Specs for video and social content to create off-platform. Each item lists the "
         "AI/search query it targets, the keywords to hit, the length and format, the hook, "
         "and the call to action — everything a producer needs to make it.", italic=True)
    if vids:
        heading("Videos to Produce", size=12)
        for r in vids:
            _render_brief(body, bullet, r)
    if socs:
        heading("Social Posts to Produce", size=12)
        for r in socs:
            _render_brief(body, bullet, r)


def generate(business_id: int) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches

    d = _load(business_id)
    b, gap, plan = d["business"], d["gap"], d["plan"]
    series = d["series"]

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    heading, body, bullet = _doc_helpers(doc)

    t = doc.add_paragraph()
    tr = t.add_run("AI Visibility \u2014 Monthly Progress Report")
    tr.bold = True; tr.font.size = Pt(22); tr.font.color.rgb = RGBColor.from_string(NAVY)
    sub = doc.add_paragraph()
    sr = sub.add_run(f"{b.get('name','')}   \u2022   {date.today().strftime('%B %Y')}")
    sr.italic = True; sr.font.color.rgb = RGBColor.from_string("555555")

    # Optional SAMPLE watermark banner (REP_REPORT_WATERMARK set) -- used for demo/
    # format previews so they can never be mistaken for a real client audit.
    import os as _os
    _wm = _os.getenv("REP_REPORT_WATERMARK", "").strip()
    if _wm:
        wp = doc.add_paragraph(); wp.paragraph_format.space_before = Pt(6)
        wp.paragraph_format.space_after = Pt(6)
        wr = wp.add_run(f"\u26A0  {_wm}  \u26A0")
        wr.bold = True; wr.font.size = Pt(13); wr.font.color.rgb = RGBColor.from_string("B00020")

    heading("Executive Summary")
    n_runs = len(series)
    if n_runs >= 2:
        first, last = series[0], series[-1]
        # round to the data's precision (goal_alignment is NUMERIC(4,2)) so the
        # notability boundary is deterministic, not subject to float representation.
        ga_change = round((last["goal_alignment"] or 0) - (first["goal_alignment"] or 0), 4)
        con_change = round((last["contested_rate"] or 0) - (first["contested_rate"] or 0), 4)
        # Honesty gate: only claim a trend when the move clears a notability floor,
        # never say "improving" when goal-alignment hasn't actually risen, and don't
        # quote a contested-framing delta that is statistically negligible.
        NOTABLE = 0.02
        improving = ga_change >= NOTABLE
        slipping = ga_change <= -NOTABLE
        direction = ("improving" if improving
                     else "slipping and needs attention" if slipping
                     else "holding roughly steady")
        con_clause = (f" Mentions of contested framing have moved by {con_change:+.0%}."
                      if abs(con_change) >= NOTABLE else "")
        if not improving and not slipping and not con_clause:
            body(f"Over {n_runs} audits, how AI assistants describe {b.get('name','your business')} "
                 f"is {direction} -- it is still early to call a trend from the data so far. The plan "
                 f"below keeps building the owned content, reviews, and third-party coverage the AI "
                 f"engines increasingly draw on. Details and the month's work follow.")
        else:
            lead = ("The accurate, positive narrative is being surfaced more often."
                    if improving else
                    "Some answers have moved the wrong way; the plan below targets those gaps."
                    if slipping else
                    "The accurate story is roughly holding while we keep building coverage.")
            # Don't pair an optimistic provenance claim with a slipping trend.
            trailer = ("We're prioritizing the owned content, reviews, and coverage that move these "
                       "answers back in your favor."
                       if slipping else
                       "The pages, reviews, and third-party coverage we've published are increasingly "
                       "what the AI engines draw on when someone asks about you.")
            body(f"Over {n_runs} audits, how AI assistants describe {b.get('name','your business')} "
                 f"is {direction}. {lead}{con_clause} {trailer} Details and the month's work follow.")
    else:
        body(f"This is the baseline audit for {b.get('name','your business')}. It captures exactly "
             f"how the major AI assistants describe you today and where the accurate story is thin. "
             f"Everything from here is measured against this starting point. The plan below front-loads "
             f"the fastest wins (reviews, your Google profile, and your first owned content).")

    chart = _trend_chart(series)
    if chart:
        heading("The Trend")
        doc.add_picture(chart, width=Inches(6.2))
        body("Higher goal-alignment and owned-content lines are better; the contested-mentions "
             "line going down is better.", italic=True)
        try:
            os.unlink(chart)  # best-effort; never let temp-file cleanup fail the report
        except OSError as e:
            log.warning("could not remove temp chart %s: %s", chart, e)

    heading("Metrics This Period")
    if n_runs >= 1:
        tbl = doc.add_table(rows=1, cols=4); tbl.style = "Light Grid Accent 1"
        for i, h in enumerate(["Metric", "First", "Latest", "Direction"]):
            tbl.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
        first = series[0]; last = series[-1]
        rows_spec = [
            ("Goal alignment (higher better)", first["goal_alignment"], last["goal_alignment"], True),
            ("Contested mentions (lower better)", first["contested_rate"], last["contested_rate"], False),
            ("Owned-content surfacing (higher better)", first["owned_rate"], last["owned_rate"], True),
        ]
        for label, fv, lv, up_good in rows_spec:
            c = tbl.add_row().cells
            c[0].text = label
            c[1].text = "n/a" if fv is None else f"{fv:.2f}"
            c[2].text = "n/a" if lv is None else f"{lv:.2f}"
            if fv is None or lv is None:
                c[3].text = "\u2014"
            else:
                delta = lv - fv
                improving = (delta > 0) if up_good else (delta < 0)
                c[3].text = ("\u2191 improving" if improving and abs(delta) >= 0.01
                             else ("\u2193 watch" if abs(delta) >= 0.01 else "\u2192 stable"))

    ba = d["before_after"]
    if ba:
        heading("How the AI Answers Changed")
        body("Real examples of how assistants answered key questions at baseline versus now:")
        for pair in ba[:3]:
            via = f"  (same engine: {pair['engine']})" if pair.get("engine") else ""
            qp = doc.add_paragraph(); qr = qp.add_run(f"Q: \u201C{pair['prompt']}\u201D{via}")
            qr.bold = True; qr.font.color.rgb = RGBColor.from_string(NAVY)
            before_txt = (pair["before"]["answer_text"] or "")[:280]
            after_txt = (pair["after"]["answer_text"] or "")[:280]
            bp = doc.add_paragraph(); bp.add_run("Before: ").bold = True
            br = bp.add_run(before_txt + ("\u2026" if len(before_txt) == 280 else ""))
            br.italic = True; br.font.color.rgb = RGBColor.from_string(RUST)
            ap = doc.add_paragraph(); ap.add_run("Now: ").bold = True
            ar = ap.add_run(after_txt + ("\u2026" if len(after_txt) == 280 else ""))
            ar.italic = True; ar.font.color.rgb = RGBColor.from_string(GREEN)

    heading("Work Completed & In Progress")
    wc = d["wo_counts"]
    if wc:
        total = sum(wc.values())
        done = wc.get("done", 0) + wc.get("verified", 0)
        body(f"{done} of {total} work orders complete ({(100*done/total):.0f}%). "
             f"{d['assets_n']} assets published to date.")
        for st in ["verified", "done", "in_progress", "blocked", "pending", "skipped"]:
            if wc.get(st):
                bullet(f"{st.replace('_',' ').title()}: {wc[st]}")
    else:
        body("Work orders not yet synced into tracking. Run the tracking module's sync-plan "
             "to begin recording execution.", italic=True)

    attr = d["attribution"]
    if attr and any(a.get("assets_in_window") for a in attr):
        heading("What Preceded the Movement")
        body("Assets published in the last interval, alongside the metric changes that followed. "
             "This is a correlation \u2014 it shows what we shipped before the change, not proof of "
             "cause:", italic=True)
        seen = set()
        for a in attr:
            aw = a.get("assets_in_window")
            items = aw if isinstance(aw, list) else (json.loads(aw) if aw else [])
            for it in items:
                if it.get("id") in seen:
                    continue
                seen.add(it.get("id"))
                bullet(f"{it.get('type','asset')}: {it.get('title','')} "
                       f"({it.get('surface','')}, {str(it.get('published_at',''))[:10]})")
        deltas = ", ".join(f"{a['metric']} {float(a['delta']):+.2f}"
                           for a in attr if a.get("delta") is not None)
        if deltas:
            body(f"Metric changes over the same window: {deltas}.")

    _section_timeline(heading, body, business_id)
    _section_acceleration(heading, body, business_id)
    _section_share_of_voice(heading, body, business_id)
    _section_per_engine(heading, body, bullet, business_id)
    _section_competitor(heading, body, bullet, business_id)

    wos = plan.get("work_orders", []) or []
    order_ids = gap.get("priority_order", []) or [w.get("wo_id") for w in wos[:6]]
    wo_by_id = {w.get("wo_id"): w for w in wos}
    shown = 0
    for oid in order_ids:
        w = wo_by_id.get(oid)
        if w:
            bullet(f"{w.get('title','')} \u2014 {w.get('phase','')} (target {w.get('target_date','')})")
            shown += 1
        if shown >= 6:
            break
    if shown == 0:
        for w in wos[:6]:
            bullet(w.get("title", ""))

    _section_root_cause_and_incidents(heading, body, bullet, business_id)
    _section_production_briefs(heading, body, bullet, business_id)

    heading("How These Results Are Achieved", size=12, color=GOLD)
    body("This program works by out-producing and out-corroborating accurate, positive content so "
         "it dominates what AI assistants surface \u2014 not by suppressing or hiding legitimate "
         "third-party views, which cannot be reliably done. Progress is re-audited each month against "
         "the baseline.", italic=True)
    if d["month_cost"]:
        body(f"[Internal note \u2014 remove for client: estimated model/API cost this month "
             f"${d['month_cost']:.2f}.]", italic=True, color="999999")

    # page footer watermark (every page) when in sample mode
    if _wm:
        try:
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            footer = doc.sections[0].footer
            fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fr = fp.add_run(_wm)
            fr.bold = True; fr.font.size = Pt(9); fr.font.color.rgb = RGBColor.from_string("B00020")
        except Exception:  # noqa: BLE001
            pass

    return _save_report(doc, b)


def _fix_settings_zoom(path: str) -> None:
    """python-docx emits a <w:zoom> with no percent attribute, which trips strict
    OOXML validators (harmless in Word). Add the attribute so the file is clean."""
    import zipfile, shutil, re as _re, tempfile as _tf
    try:
        tmp = _tf.NamedTemporaryFile(suffix=".docx", delete=False).name
        with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.namelist():
                data = zin.read(item)
                if item == "word/settings.xml":
                    txt = data.decode("utf-8")
                    if "w:zoom" in txt and "w:percent" not in txt:
                        txt = _re.sub(r"<w:zoom\s*/>", '<w:zoom w:percent="100"/>', txt)
                        txt = _re.sub(r"<w:zoom(?![^>]*w:percent)([^>]*)/>",
                                      r'<w:zoom\1 w:percent="100"/>', txt)
                        data = txt.encode("utf-8")
                zout.writestr(item, data)
        shutil.move(tmp, path)
    except Exception as e:  # noqa: BLE001  -- never fail the report over a cosmetic fix
        log.warning("settings.xml zoom fix skipped: %s", e)


def main() -> None:
    ap = argparse.ArgumentParser(description="Client report generator")
    sub = ap.add_subparsers(dest="cmd", required=True)
    pm = sub.add_parser("monthly"); pm.add_argument("--business-id", type=int, required=True)
    args = ap.parse_args()
    if args.cmd == "monthly":
        generate(args.business_id)


if __name__ == "__main__":
    main()
