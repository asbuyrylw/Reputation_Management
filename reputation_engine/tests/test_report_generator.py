"""Smoke test for report_generator.generate -- requires REP_TEST_DSN.

The report hub is otherwise untested. generate() makes NO LLM calls (the
narrative is templated Python; the optional timeline/acceleration/citation/
competitor sections are DB-only and individually try/except-wrapped inside
generate()), so this needs no LLM monkeypatching. matplotlib Agg is selected
inside report_generator._trend_chart() itself (matplotlib.use("Agg")), so no
headless setup is required here. We only redirect OUTPUT_DIR to a tmp dir so the
generated .docx is isolated and auto-cleaned by pytest's tmp_path.
"""

from __future__ import annotations

import os

from conftest import requires_db


def _seed_business(conn, name="Reportco LLC", contested="MLM,scam"):
    row = conn.execute(
        "INSERT INTO businesses (name, domain, services, goal, contested_terms, geo) "
        "VALUES (%s,%s,%s,%s,%s,%s) RETURNING id",
        (name, "reportco.com", "insurance", "win local queries", contested, "Cincinnati OH"),
    ).fetchone()
    conn.commit()
    return row["id"]


def _complete_run(conn, bid, days_ago):
    """A finished+complete audit_run, dated `days_ago` in the past."""
    return conn.execute(
        "INSERT INTO audit_runs (business_id, finished_at, status) "
        "VALUES (%s, now() - make_interval(days => %s), 'complete') RETURNING id",
        (bid, days_ago),
    ).fetchone()["id"]


def _answer(conn, rid, bid, prompt, text, ga, contested, owned):
    conn.execute(
        "INSERT INTO answers (run_id,business_id,engine,prompt,answer_text,"
        "goal_alignment,mentions_contested,surfaces_owned,failed) "
        "VALUES (%s,%s,'e',%s,%s,%s,%s,%s,false)",
        (rid, bid, prompt, text, ga, contested, owned),
    )


def _doc_text(path):
    from docx import Document
    return "\n".join(p.text for p in Document(path).paragraphs)


@requires_db
def test_generate_produces_nonempty_docx(fresh_schema, tmp_path, monkeypatch):
    conn = fresh_schema
    from rep_engine import report_generator as rg

    bid = _seed_business(conn)

    # Two COMPLETE runs sharing one prompt: baseline (weak) then improved (strong).
    # Two finished runs light up the trend chart, metrics table, and before/after
    # sections -- the most code-heavy branches of generate().
    prompt = "Is Reportco LLC trustworthy?"
    r1 = _complete_run(conn, bid, days_ago=30)
    r2 = _complete_run(conn, bid, days_ago=0)
    _answer(conn, r1, bid, prompt, "Some sources call Reportco an MLM.", 0.10, True, False)
    _answer(conn, r2, bid, prompt, "Reportco LLC is a trusted local insurance firm.", 0.80, False, True)
    conn.commit()

    # Isolate output so the .docx (and its temp chart PNG) never touch the repo's
    # ./output dir; tmp_path is auto-removed by pytest, so no manual cleanup.
    monkeypatch.setattr(rg, "OUTPUT_DIR", str(tmp_path))
    # Ensure no demo watermark branch is taken.
    monkeypatch.delenv("REP_REPORT_WATERMARK", raising=False)

    path = rg.generate(bid)

    # File was produced, lives under our tmp dir, and is non-empty.
    assert os.path.isfile(path)
    assert os.path.dirname(os.path.abspath(path)) == os.path.abspath(str(tmp_path))
    assert os.path.getsize(path) > 0

    # Reopen the .docx and confirm key narrative content rendered.
    from docx import Document
    full_text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "Reportco LLC" in full_text
    assert "Executive Summary" in full_text


@requires_db
def test_before_after_compares_same_engine(fresh_schema):
    """Before/Now must compare the SAME engine -- not the max-goal_alignment answer
    per run (which could put Perplexity 'Before' against ChatGPT 'Now')."""
    conn = fresh_schema
    from rep_engine import report_generator as rg
    bid = _seed_business(conn)
    p = "Is Reportco LLC trustworthy?"
    r1 = _complete_run(conn, bid, days_ago=30)
    r2 = _complete_run(conn, bid, days_ago=0)
    # run1: engineA weak (0.2), engineB strong (0.6); run2: engineA strong (0.9),
    # engineB weak (0.3). A naive max-per-run would pair engineB(before) vs engineA(now).
    for rid, eng, ga in [(r1, "engineA", 0.2), (r1, "engineB", 0.6),
                         (r2, "engineA", 0.9), (r2, "engineB", 0.3)]:
        conn.execute("INSERT INTO answers (run_id,business_id,engine,prompt,answer_text,"
                     "goal_alignment,failed) VALUES (%s,%s,%s,%s,'ans',%s,false)",
                     (rid, bid, eng, p, ga))
    conn.commit()
    pairs = rg._before_after(conn, bid)
    assert pairs, "expected a before/after pair"
    pair = pairs[0]
    assert pair["before"]["engine"] == pair["after"]["engine"]   # honest like-for-like
    assert pair["engine"] in {"engineA", "engineB"}


@requires_db
def test_exec_summary_does_not_overclaim_on_thin_change(fresh_schema, tmp_path, monkeypatch):
    """With a sub-threshold change, the exec summary must not claim 'improving' --
    it should hedge that it is too early to call a trend."""
    conn = fresh_schema
    from rep_engine import report_generator as rg
    bid = _seed_business(conn)
    p = "Is Reportco LLC trustworthy?"
    r1 = _complete_run(conn, bid, days_ago=30)
    r2 = _complete_run(conn, bid, days_ago=0)
    # near-identical metrics (ga 0.40 -> 0.41 = +0.01, below the 0.02 floor)
    _answer(conn, r1, bid, p, "ans baseline", 0.40, False, True)
    _answer(conn, r2, bid, p, "ans now", 0.41, False, True)
    conn.commit()
    monkeypatch.setattr(rg, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.delenv("REP_REPORT_WATERMARK", raising=False)
    path = rg.generate(bid)
    text = _doc_text(path)
    assert "still early to call a trend" in text
    assert "is improving." not in text                           # no over-claim


@requires_db
def test_exec_summary_slipping_uses_honest_tone(fresh_schema, tmp_path, monkeypatch):
    """A declining trend must read as slipping/needs-attention, never paired with an
    optimistic 'increasingly what the AI engines draw on' provenance claim."""
    conn = fresh_schema
    from rep_engine import report_generator as rg
    bid = _seed_business(conn)
    p = "Is Reportco LLC trustworthy?"
    r1 = _complete_run(conn, bid, days_ago=30)
    r2 = _complete_run(conn, bid, days_ago=0)
    _answer(conn, r1, bid, p, "ans baseline", 0.60, False, True)
    _answer(conn, r2, bid, p, "ans now", 0.40, False, True)   # ga -0.20 -> slipping
    conn.commit()
    monkeypatch.setattr(rg, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.delenv("REP_REPORT_WATERMARK", raising=False)
    text = _doc_text(rg.generate(bid))
    assert "slipping and needs attention" in text
    assert "moved the wrong way" in text
    assert "move these answers back in your favor" in text     # honest forward-looking trailer
    assert "is improving." not in text


@requires_db
def test_exec_summary_reports_contested_move_when_alignment_holds(fresh_schema, tmp_path, monkeypatch):
    """goal_alignment flat but contested framing jumps: the else-branch must fire,
    report the contested move, and NOT hedge with 'still early to call a trend'."""
    conn = fresh_schema
    from rep_engine import report_generator as rg
    bid = _seed_business(conn)
    p = "Is Reportco LLC trustworthy?"
    r1 = _complete_run(conn, bid, days_ago=30)
    r2 = _complete_run(conn, bid, days_ago=0)
    _answer(conn, r1, bid, p, "ans baseline", 0.50, False, True)
    _answer(conn, r2, bid, p, "ans now", 0.50, True, True)    # ga flat, contested 0->100%
    conn.commit()
    monkeypatch.setattr(rg, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.delenv("REP_REPORT_WATERMARK", raising=False)
    text = _doc_text(rg.generate(bid))
    assert "holding roughly steady" in text
    assert "contested framing have moved by" in text           # con_clause fired (else branch)
    assert "still early to call a trend" not in text


@requires_db
def test_report_surfaces_root_cause_and_incidents(fresh_schema, tmp_path, monkeypatch):
    """When the agentic layer has run, the report shows the root-cause + flagged
    incidents; otherwise the section is a no-op (verified by the smoke test)."""
    import json
    conn = fresh_schema
    from rep_engine import report_generator as rg
    bid = _seed_business(conn)
    r1 = _complete_run(conn, bid, days_ago=30)
    r2 = _complete_run(conn, bid, days_ago=0)
    _answer(conn, r1, bid, "Is Reportco trustworthy?", "Some call it an MLM.", 0.10, True, False)
    _answer(conn, r2, bid, "Is Reportco trustworthy?", "Trusted local firm.", 0.80, False, True)
    conn.execute("INSERT INTO root_cause (business_id, model) VALUES (%s,%s)",
                 (bid, json.dumps({"summary": "The MLM narrative is driven by ripoffreport.com",
                                   "primary_sources": [{"url": "https://ripoffreport.com/x",
                                                        "why": "top contested citation"}],
                                   "recommended_counters": ["3 owned blog posts"]})))
    conn.execute("INSERT INTO incidents (business_id, mention_url, severity, status) "
                 "VALUES (%s,'https://r/1','high','pending_human_review')", (bid,))
    conn.commit()
    monkeypatch.setattr(rg, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.delenv("REP_REPORT_WATERMARK", raising=False)
    text = _doc_text(rg.generate(bid))
    assert "Why the Contested Narrative Surfaces" in text
    assert "ripoffreport.com" in text
    assert "New Contested Mentions Flagged This Period" in text
